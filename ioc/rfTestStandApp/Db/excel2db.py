#!/usr/bin/env python3

from openpyxl import load_workbook
import sys

# Path to your Excel file
file_path = sys.argv[1]

# Load the workbook
workbook = load_workbook(filename=file_path)

# support the following record types
channel_types = ["aai", "aao", "ai", "ao", "aSub", "bi", "bo", "calcout", "calc", "compress", "dfanout", "event", "fanout", "histogram", "longin", "longout", "lsi", "lso", "mbbiDirect", "mbbi", "mbboDirect", "mbbo", "permissive", "prinf", "sel", "seq", "state", "stringin", "stringout", "subArray", "sub", "waveform"]
header_line = []

def generate_comment(row):
    result = "# "
    for cell in row:
        if cell is not None:
            result = result + " " + str(cell)
    return result + "\n"

def parse_row(row):
    first_field = row[0]
    if first_field in channel_types:
        return generate_channel(row)
    else:
        return generate_comment(row)

def generate_channel(row):
    result = ""
    for index, field_name in enumerate(header_line):
        field_value = row[index]
        if field_value is not None:
            if field_name == "recordType":
                result = result + "record(" + field_value + ", "
            elif field_name == "recordName":
                result = result + '"' + field_value + '") {\n'
            elif field_name == "DESC":
                field_value_trimmed = field_value
                if len(field_value) > 39:
                    field_value_trimmed = field_value[0:39]
                    result = result + f'    # {field_value}\n'
                result = result + f'    field({field_name}, "{field_value_trimmed}")\n'
            elif field_name == "Comment":
                result = result + f'    # {field_value}\n'
            else:
                result = result + f'    field({field_name}, "{field_value}")\n'
        else:
            pass
    result = result + "}\n\n"
    return result

for sheet_name in workbook.sheetnames:
    sheet = workbook[sheet_name]
    result = ""
    for index, row in enumerate(sheet.iter_rows(values_only=True)):
        # print(index)
        if index == 0:
            header_line = row
        else:
            result = result + parse_row(row)

    # each sheet is a file
    file_name = file_path.replace(".xlsx", "_") + sheet_name + ".db"
    with open(file_name, "w") as f:
        f.write(result + "\n")
        print("Success:\n  ", file_path, "-->", file_name, "\n")


